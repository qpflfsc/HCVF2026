# -*- coding: utf-8 -*-
"""
QPFL HCVF Stakeholder Consultation - Excel Auto Generator
Author = "Lộc Vũ Trung"
"""
import os
import sys
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_JSON = os.path.join(BASE_DIR, 'data', 'data.json')
EXCEL_PATH = os.path.join(BASE_DIR, 'QPFL_Tong_Hop_Ket_Qua_Tham_Van_2026.xlsx')

def generate_excel():
    if not os.path.exists(DATA_JSON):
        print(f"Error: {DATA_JSON} not found!")
        return

    with open(DATA_JSON, 'r', encoding='utf-8') as f:
        data_store = json.load(f)

    records = data_store.get('records', [])
    updated_at = data_store.get('updated_at', datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S'))

    wb = openpyxl.Workbook()
    wb.properties.creator = "Lộc Vũ Trung"
    wb.properties.title = "QPFL FSC HCVF 2026 - Tổng Hợp Tham Vấn"

    font_title = Font(name='Times New Roman', size=14, bold=True, color='1B365D')
    font_sub = Font(name='Times New Roman', size=10, italic=True)
    font_head = Font(name='Times New Roman', size=10, bold=True, color='FFFFFF')
    font_bold = Font(name='Times New Roman', size=10, bold=True)
    font_regular = Font(name='Times New Roman', size=10)

    fill_navy = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    fill_alt = PatternFill(start_color='F9FAFC', end_color='F9FAFC', fill_type='solid')

    thin = Side(border_style='thin', color='D9D9D9')
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 1. Sheet CHI_TIET_PHIEU_THAM_VAN
    ws1 = wb.active
    ws1.title = "CHI_TIET_PHIEU_THAM_VAN"

    ws1['A1'] = "CÔNG TY TNHH TRỒNG RỪNG QUY NHƠN (QPFL) — CHỨNG CHỈ FSC"
    ws1['A1'].font = font_title
    ws1['A2'] = "BẢNG KẾT QUẢ THAM VẤN CÁC BÊN LIÊN QUAN VỀ RỪNG CÓ GIÁ TRỊ BẢO TỒN CAO (HCVF) NĂM 2026"
    ws1['A2'].font = Font(name='Times New Roman', size=12, bold=True, color='2E75B6')
    ws1['A3'] = f"Hệ thống thu thập trực tuyến qpflfsc.github.io/HCVF2026 | Cập nhật: {updated_at} | Tác giả: Lộc Vũ Trung"
    ws1['A3'].font = font_sub

    headers1 = [
        "STT", "Thời gian ghi nhận", "Họ và tên người đại diện", "Chức vụ / Học vị",
        "Cơ quan / Đơn vị / Trường Viện / Thôn xóm", "Số điện thoại / Zalo", "Email người phản hồi",
        "Địa bàn xã / khu vực liên quan", "Tư cách tham vấn",
        "1. Đa dạng ĐV-TV & Lâm sản phụ (Chọn)", "1. Ghi chú ý kiến loài",
        "2. Nguồn nước & Chống xói mòn HCV 4 (Chọn)", "2. Ghi chú suối/hầm lắng",
        "3. Mồ mả, PCCC & Ranh giới (Chọn)", "3. Ghi chú ranh giới/mồ mả/PCCC",
        "4. Giải pháp & Cam kết QPFL (Chọn)", "4. Đề xuất giải pháp thêm",
        "Kiến nghị / Đề xuất chung", "Người gửi / xác nhận", "Cán bộ tiếp nhận QPFL"
    ]

    for col_idx, h in enumerate(headers1, start=1):
        c = ws1.cell(5, col_idx, h)
        c.font = font_head
        c.fill = fill_navy
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border_all
    ws1.row_dimensions[5].height = 28

    curr_r = 6
    for rec in records:
        row_vals = [
            rec.get('id', curr_r - 5),
            rec.get('timestamp', ''),
            rec.get('rep_name', ''),
            rec.get('rep_title', ''),
            rec.get('org_name', ''),
            rec.get('rep_phone', ''),
            rec.get('rep_email', ''),
            rec.get('commune', ''),
            rec.get('stk_group', ''),
            rec.get('q1_bio', ''),
            rec.get('q1_note', ''),
            rec.get('q2_water', ''),
            rec.get('q2_note', ''),
            rec.get('q3_social', ''),
            rec.get('q3_note', ''),
            rec.get('q4_commit', ''),
            rec.get('q4_note', ''),
            rec.get('feedback', ''),
            rec.get('signer_name', rec.get('rep_name', '')),
            rec.get('receiver_name', 'Trần Quang Tâm (Phòng Kỹ Thuật QPFL)')
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(curr_r, col_idx, val)
            cell.font = font_regular
            cell.border = border_all
            if col_idx in [1, 2, 6, 7]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            if (curr_r - 5) % 2 == 0:
                cell.fill = fill_alt
        curr_r += 1

    # 2. Sheet TONG_HOP_THONG_KE (100% Natural Formulas)
    ws2 = wb.create_sheet(title="TONG_HOP_THONG_KE")
    ws2['A1'] = "TỔNG HỢP VÀ THỐNG KÊ TỶ LỆ ĐỒNG THUẬN CỦA CÁC BÊN LIÊN QUAN (HCV 1, 4, 6)"
    ws2['A1'].font = font_title
    ws2['A2'] = "Dữ liệu phục vụ Báo cáo Tham vấn và Bằng chứng đóng Major CAR 03"
    ws2['A2'].font = Font(name='Times New Roman', size=11, bold=True, color='2E75B6')
    ws2['A3'] = "Tác giả: Lộc Vũ Trung | Công thức Excel tự động tổng hợp theo thời gian thực"
    ws2['A3'].font = font_sub

    ws2['A5'] = "TỔNG SỐ PHIẾU ĐÃ THAM VẤN:"
    ws2['A5'].font = font_bold
    ws2['C5'] = "=COUNTA('CHI_TIET_PHIEU_THAM_VAN'!A6:A1000)"
    ws2['C5'].font = Font(name='Times New Roman', size=14, bold=True, color='1B365D')

    ws2['A7'] = "1. CƠ CẤU NHÓM BÊN LIÊN QUAN THAM GIA THAM VẤN"
    ws2['A7'].font = font_bold

    headers_stk = ["STT", "Nhóm Bên Liên Quan", "Số Lượng (Phiếu)", "Tỷ Lệ (%)"]
    for ci, h in enumerate(headers_stk, start=1):
        c = ws2.cell(8, ci, h)
        c.font = font_head
        c.fill = fill_navy
        c.alignment = Alignment(horizontal='center')
        c.border = border_all

    groups = [
        ("Cơ quan Quản lý Lâm nghiệp (Kiểm lâm)", "Kiểm lâm"),
        ("Chính quyền địa phương (Trưởng thôn / Xã)", "thôn"),
        ("Doanh nghiệp lâm nghiệp / Đối tác giáp ranh", "Đối tác"),
        ("Cộng đồng dân cư / Người dân địa phương", "Người dân"),
        ("Chuyên gia / Nhà khoa học / Giảng viên", "Chuyên gia"),
        ("Công nhân / Bảo vệ rừng", "Công nhân")
    ]

    r_idx = 9
    for idx, (g_name, g_val) in enumerate(groups, start=1):
        ws2.cell(r_idx, 1, idx).alignment = Alignment(horizontal='center')
        ws2.cell(r_idx, 2, g_name).alignment = Alignment(horizontal='left')
        ws2.cell(r_idx, 3, f"=COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!I$6:I$1000, \"*{g_val}*\")").number_format = '#,##0'
        ws2.cell(r_idx, 4, f"=IF($C$5>0, C{r_idx}/$C$5, 0)").number_format = '0.0%'
        for ci in range(1, 5):
            ws2.cell(r_idx, ci).font = font_regular
            ws2.cell(r_idx, ci).border = border_all
        r_idx += 1

    r_idx += 1
    ws2.cell(r_idx, 1, "2. CHỈ SỐ ĐỒNG THUẬN CÁC GIẢI PHÁP ĐÓNG CAR 03 (HCV 1, HCV 4, HCV 6)").font = font_bold
    r_idx += 1

    headers_ind = ["STT", "Tiêu Chí & Giá Trị Bảo Tồn", "Mô Tả Tiêu Chí Đánh Giá", "Số Lượng Đồng Thuận", "Tỷ Lệ Đồng Thuận (%)"]
    for ci, h in enumerate(headers_ind, start=1):
        c = ws2.cell(r_idx, ci, h)
        c.font = font_head
        c.fill = fill_navy
        c.alignment = Alignment(horizontal='center')
        c.border = border_all
    r_idx += 1

    indicators = [
        ("Đa dạng sinh học & Lâm sản phụ [HCV 1]", "Nhất trí bảo tồn loài và quy chế thu hái lâm sản phụ", "=COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!J$6:J$1000, \"*nhất trí*\") + COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!J$6:J$1000, \"*Mức 1*\")"),
        ("Nguồn nước & Chống xói mòn [HCV 4]", "An tâm với 63 hầm lắng và dải cây giữ đất giữ nước", "=COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!L$6:L$1000, \"*an tâm*\") + COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!L$6:L$1000, \"*Mức 1*\")"),
        ("Mồ mả, Phong tục & An toàn PCCC [HCV 6]", "Hài lòng về đai đệm bảo vệ mồ mả và ổn định ranh giới", "=COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!N$6:N$1000, \"*hài lòng*\") + COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!N$6:N$1000, \"*Mức 1*\")"),
        ("Cam kết quản lý rừng bền vững QPFL [Chung]", "Đồng thuận và ủng hộ cao công ty duy trì FSC", "=COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!P$6:P$1000, \"*ủng hộ*\") + COUNTIF('CHI_TIET_PHIEU_THAM_VAN'!P$6:P$1000, \"*Mức 1*\")")
    ]

    for idx, (title, desc, formula) in enumerate(indicators, start=1):
        ws2.cell(r_idx, 1, idx).alignment = Alignment(horizontal='center')
        ws2.cell(r_idx, 2, title).alignment = Alignment(horizontal='left')
        ws2.cell(r_idx, 3, desc).alignment = Alignment(horizontal='left')
        ws2.cell(r_idx, 4, formula).number_format = '#,##0'
        ws2.cell(r_idx, 5, f"=IF($C$5>0, D{r_idx}/$C$5, 0)").number_format = '0.0%'
        for ci in range(1, 6):
            ws2.cell(r_idx, ci).font = font_regular
            ws2.cell(r_idx, ci).border = border_all
        r_idx += 1

    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if not val.startswith('='):
                    max_len = max(max_len, len(val))
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 12)

    ws1.column_dimensions['C'].width = 25
    ws1.column_dimensions['D'].width = 20
    ws1.column_dimensions['E'].width = 35
    ws2.column_dimensions['B'].width = 40
    ws2.column_dimensions['C'].width = 45

    wb.save(EXCEL_PATH)
    print(f"Successfully updated {EXCEL_PATH} with {len(records)} records. Author = Lộc Vũ Trung")

if __name__ == '__main__':
    generate_excel()
